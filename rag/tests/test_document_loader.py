from pathlib import Path

import docx
import pytest

from src.document_loader import (
    DocumentLoadError,
    EmptyDocumentError,
    LoadStats,
    iter_documents,
    load_directory,
    load_file,
    project_from_path,
    slugify_project_id,
)
from tests.conftest import write_minimal_pdf


def test_slugify():
    assert slugify_project_id("Project Alpha") == "project_alpha"
    assert slugify_project_id("  ") == "unassigned"


def test_project_from_path_levels():
    assert project_from_path(Path("pdfs/project_alpha/a.pdf"), 2) == "project_alpha"
    assert project_from_path(Path("pdfs/a.pdf"), 2) == "unassigned"
    assert project_from_path(Path("project_x/a.pdf"), 1) == "project_x"
    assert project_from_path(Path("pdfs/project_alpha/a.pdf"), 0) == "unassigned"


def test_load_directory_all_types(docs_dir):
    docs, stats = load_directory(docs_dir)
    by_name = {d.source_file: d for d in docs}
    assert set(by_name) == {"risk_register.csv", "status.txt", "plan.pdf"}
    assert stats.documents_loaded == 3 and stats.files_failed == 0
    assert by_name["status.txt"].project_id == "project_alpha"
    assert by_name["status.txt"].source_path == "text/project_alpha/status.txt"
    assert by_name["plan.pdf"].project_id == "project_beta"
    assert [s.page_number for s in by_name["plan.pdf"].sections] == [1, 2]
    assert len(by_name["risk_register.csv"].doc_hash) == 64


def test_csv_rows_headers_and_project_column(docs_dir):
    doc = next(d for d in load_directory(docs_dir)[0] if d.doc_type == "csv")
    assert doc.columns[:3] == ["Risk ID", "Project", "Risk Description"]
    assert [s.row_number for s in doc.sections] == [1, 2, 3]
    assert [s.project_id for s in doc.sections] == ["project_alpha", "project_alpha", "project_beta"]
    assert "Mitigation Plan: Hire two contractors" in doc.sections[0].text


def test_csv_ragged_blank_and_duplicate_headers(tmp_path):
    f = tmp_path / "r.csv"
    f.write_text("Name,,Name\nA,x,y,extra\n\n,,\nB\n", encoding="utf-8")
    doc = load_file(f, tmp_path)
    assert doc.columns == ["Name", "column_2", "Name_2"]
    assert len(doc.sections) == 2
    assert "extra_4: extra" in doc.sections[0].text


def test_txt_encodings(tmp_path):
    (tmp_path / "a.txt").write_bytes("caf\xe9 risk".encode("cp1252"))
    assert "risk" in load_file(tmp_path / "a.txt", tmp_path).sections[0].text


def test_docx_headings_and_tables(tmp_path):
    d = docx.Document()
    d.add_heading("Risks", level=1)
    d.add_paragraph("Resource shortage is likely.")
    table = d.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Risk", "Mitigation"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Staff loss", "Hire contractors"
    d.add_heading("Budget", level=1)
    d.add_paragraph("Budget is reviewed monthly.")
    d.save(tmp_path / "plan.docx")
    doc = load_file(tmp_path / "plan.docx", tmp_path)
    assert [s.section for s in doc.sections] == ["Risks", "Budget"]
    assert "Risk: Staff loss; Mitigation: Hire contractors" in doc.sections[0].text


def test_corrupted_files_are_reported_not_raised(tmp_path):
    (tmp_path / "bad.pdf").write_bytes(b"this is not a pdf")
    (tmp_path / "bad.docx").write_bytes(b"not a zip")
    (tmp_path / "bin.txt").write_bytes(b"\x00\x01\x02binary")
    (tmp_path / "empty.txt").write_text("   \n")
    (tmp_path / "data.zip").write_bytes(b"PK")
    (tmp_path / "ok.txt").write_text("fine")
    stats = LoadStats()
    docs = list(iter_documents(tmp_path, stats))
    assert [d.source_file for d in docs] == ["ok.txt"]
    assert {p for p, _ in stats.failures} == {"bad.pdf", "bad.docx", "bin.txt"}
    assert stats.empty_files == ["empty.txt"]
    assert stats.unsupported_files == ["data.zip"]
    assert stats.files_found == 6


def test_blank_pdf_is_empty(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    with open(tmp_path / "scan.pdf", "wb") as fh:
        writer.write(fh)
    with pytest.raises(EmptyDocumentError, match="OCR"):
        load_file(tmp_path / "scan.pdf", tmp_path)


def test_hidden_and_lock_files_ignored_and_recursion(tmp_path):
    (tmp_path / "a" / "b").mkdir(parents=True)
    (tmp_path / "a" / "b" / "deep.txt").write_text("deep")
    (tmp_path / "~$lock.docx").write_text("x")
    (tmp_path / ".hidden.txt").write_text("x")
    docs, stats = load_directory(tmp_path)
    assert [d.source_path for d in docs] == ["a/b/deep.txt"]
    assert load_directory(tmp_path, recursive=False)[0] == []


def test_unsupported_extension_and_missing_dir(tmp_path):
    (tmp_path / "x.xyz").write_text("x")
    with pytest.raises(DocumentLoadError):
        load_file(tmp_path / "x.xyz", tmp_path)
    with pytest.raises(FileNotFoundError):
        list(iter_documents(tmp_path / "nope"))


def test_file_size_limit(tmp_path):
    (tmp_path / "big.txt").write_text("x" * 2_000_000)
    with pytest.raises(DocumentLoadError, match="limit"):
        load_file(tmp_path / "big.txt", tmp_path, max_file_mb=1)
